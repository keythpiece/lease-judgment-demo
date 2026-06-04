from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .display import display_text, display_value, review_text
from .utils import flatten_evidence, load_yaml


def write_excel_assessment(
    result: dict[str, Any],
    output_path: str | Path,
    template_path: str | Path | None = None,
    mapping_path: str | Path | None = None,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    base = Path(__file__).resolve().parents[1]
    mapping = load_yaml(mapping_path or base / "config" / "cell_mapping.yaml")

    if template_path and Path(template_path).exists():
        wb = load_workbook(template_path)
    else:
        wb = Workbook()
    ws = wb.active
    ws.title = ws.title if ws.title != "Sheet" else "リース判定"

    _write_default_headers(ws)
    _write_mapped_cells(ws, mapping, result)
    _write_evidence_sheet(wb, result)
    wb.save(output_path)
    return output_path


def _write_default_headers(ws) -> None:
    fill = PatternFill("solid", fgColor="D9EAF7")
    headers = [
        ("A1", "リース会計 判定結果"),
        ("A2", "リース対象判定"),
        ("A3", "リース分類"),
        ("A4", "信頼度"),
        ("A5", "契約書名"),
        ("A6", "契約相手先"),
        ("A7", "契約開始日"),
        ("A8", "契約終了日"),
        ("A9", "契約期間（月）"),
        ("A12", "Step1結果"),
        ("E19", "Step2項目"),
        ("F19", "判定"),
        ("G19", "コメント"),
        ("B15", "サマリー理由"),
    ]
    for cell, value in headers:
        ws[cell] = value
        ws[cell].font = Font(bold=True)
        ws[cell].fill = fill
    step2_labels = [
        ("E20", "形式判定"),
        ("E21", "取引アイテム"),
        ("E22", "資産の特定"),
        ("E23", "入替権"),
        ("E24", "経済的便益"),
        ("E25", "指図権"),
        ("E26", "稼働権・設計関与"),
        ("E27", "資産種別"),
        ("E29", "Step2結果"),
        ("B10", "リース開始日"),
        ("B11", "リース終了日"),
        ("B12", "リース期間（月）"),
        ("B20", "経済的耐用年数（月）"),
        ("B30", "原資産の公正価値"),
        ("E35", "75%テスト"),
        ("E36", "90%PVテスト"),
        ("E38", "分類結果"),
    ]
    for cell, value in step2_labels:
        ws[cell] = value
        ws[cell].font = Font(bold=True)
    for col in range(1, 8):
        ws.column_dimensions[get_column_letter(col)].width = 22 if col != 7 else 55
    ws["B15"].alignment = Alignment(wrap_text=True, vertical="top")


def _write_mapped_cells(ws, mapping: dict[str, Any], result: dict[str, Any]) -> None:
    summary = result.get("contract_summary", {})
    step1 = result.get("step1", {})
    step2 = result.get("step2", {})
    step3 = result.get("step3", {})
    final = result.get("final_result", {})
    guidance_by_key = {
        item.get("key"): item for item in final.get("human_review_guidance", []) or []
    }
    values = {
        "step1": {
            "contract_title": summary.get("contract_title", ""),
            "counterparty": summary.get("counterparty", ""),
            "start_date": summary.get("start_date", ""),
            "end_date": summary.get("end_date", ""),
            "contract_term_months": summary.get("contract_term_months"),
            "step1_result": display_value(step1.get("step1_result", "")),
        },
        "step2": {
            "formal_assessment_answer": display_value(step2.get("formal_assessment", {}).get("answer")),
            "formal_assessment_comment": _comment_text(step2.get("formal_assessment", {})),
            "asset_item_answer": display_value(step2.get("asset_item_assessment", {}).get("answer")),
            "asset_item_comment": _comment_text(step2.get("asset_item_assessment", {}), guidance_by_key.get("asset_item_assessment")),
            "identified_asset_answer": display_value(step2.get("identified_asset", {}).get("answer")),
            "identified_asset_comment": _comment_text(step2.get("identified_asset", {}), guidance_by_key.get("identified_asset")),
            "substitution_right_answer": display_value(step2.get("substitution_right", {}).get("answer")),
            "substitution_right_comment": _comment_text(step2.get("substitution_right", {}), guidance_by_key.get("substitution_right")),
            "economic_benefits_answer": display_value(step2.get("economic_benefits", {}).get("answer")),
            "economic_benefits_comment": _comment_text(step2.get("economic_benefits", {}), guidance_by_key.get("economic_benefits")),
            "right_to_direct_use_answer": display_value(step2.get("right_to_direct_use", {}).get("answer")),
            "right_to_direct_use_comment": _comment_text(step2.get("right_to_direct_use", {}), guidance_by_key.get("right_to_direct_use")),
            "operation_right_answer": display_value(step2.get("operation_right_or_design_involvement", {}).get("answer")),
            "operation_right_comment": _comment_text(step2.get("operation_right_or_design_involvement", {}), guidance_by_key.get("operation_right_or_design_involvement")),
            "asset_type_answer": display_value(step2.get("asset_type", {}).get("answer")),
            "asset_type_comment": _comment_text(step2.get("asset_type", {}), guidance_by_key.get("asset_type")),
            "step2_result": display_value(step2.get("step2_result", "")),
        },
        "step3": {
            "lease_start_date": summary.get("start_date", ""),
            "lease_end_date": summary.get("end_date", ""),
            "lease_term_months": step3.get("lease_term_months"),
            "economic_life_months": step3.get("economic_life_months"),
            "fair_value": step3.get("fair_value"),
            "criterion_1_answer": display_value(step3.get("criterion_1_75_percent_test", {}).get("answer")),
            "criterion_1_comment": step3.get("criterion_1_75_percent_test", {}).get("reason"),
            "criterion_2_answer": display_value(step3.get("criterion_2_90_percent_pv_test", {}).get("answer")),
            "criterion_2_comment": step3.get("criterion_2_90_percent_pv_test", {}).get("reason"),
            "classification_result": display_value(step3.get("classification_result")),
        },
        "final_result": {
            "lease_applicability": display_value(final.get("lease_applicability")),
            "lease_classification": display_value(final.get("lease_classification")),
            "confidence": display_value(final.get("confidence")),
            "summary_reason": display_text(final.get("summary_reason")),
        },
    }
    for section, items in mapping.items():
        for key, cell in items.items():
            value = values.get(section, {}).get(key, "")
            ws[cell] = value
            ws[cell].alignment = Alignment(wrap_text=True, vertical="top")
            if key.endswith("_comment") and value:
                ws[cell].comment = Comment(str(value)[:30000], "AI")


def _write_evidence_sheet(wb, result: dict[str, Any]) -> None:
    if "AI判定根拠" in wb.sheetnames:
        del wb["AI判定根拠"]
    ws = wb.create_sheet("AI判定根拠")
    headers = ["No", "ページ", "判定項目", "根拠テキスト", "理由", "ハイライト対象"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAD3")
    for idx, ev in enumerate(flatten_evidence(result), start=1):
        ws.append(
            [
                idx,
                ev.get("page"),
                ev.get("source_item", ""),
                ev.get("text", ""),
                ev.get("why_relevant", ""),
                ev.get("highlight_required", True),
            ]
        )
    widths = [8, 10, 24, 90, 45, 16]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _comment_text(item: dict[str, Any], guidance: dict[str, Any] | None = None) -> str:
    pages = sorted({str(ev.get("page")) for ev in item.get("evidence", []) if ev.get("page")})
    page_text = f" 根拠ページ: {', '.join(pages)}" if pages else ""
    guide_text = review_text(guidance)
    guide_text = f"\n確認事項:\n{guide_text}" if guide_text else ""
    return f"{item.get('reason', '')}{page_text}{guide_text}".strip()
